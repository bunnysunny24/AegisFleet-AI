"""
Generates the official Hackathon Solution Document in Microsoft Word (.docx) format
matching the exact sections of Motorq_Hackathon_Solution_Document_Template.docx.
"""
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor


def create_solution_docx(md_path: str, output_path: str):
    doc = Document()

    # Configure Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Read markdown
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []
    in_code = False
    code_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        parsed_rows = []
        for line in tbl_lines:
            parts = [c.strip() for c in line.split('|')[1:-1]]
            # Skip separator row like |---|---|
            if parts and all(re.match(r'^:?-+:?$', c) for c in parts):
                continue
            parsed_rows.append(parts)

        if not parsed_rows:
            return

        cols_count = len(parsed_rows[0])
        table = doc.add_table(rows=len(parsed_rows), cols=cols_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        for r_idx, row_data in enumerate(parsed_rows):
            for c_idx, cell_value in enumerate(row_data):
                if c_idx < cols_count:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_value
                    # Format header
                    if r_idx == 0:
                        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
                        cell._tc.get_or_add_tcPr().append(shading_elm)
                        for p in cell.paragraphs:
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            for run in p.runs:
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                run.font.size = Pt(9.5)
                    else:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(9.5)
        doc.add_paragraph()

    for line in lines:
        stripped = line.strip()

        # Handle Code Blocks
        if stripped.startswith("```"):
            if in_code:
                # Flush code block
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
                p._p.get_or_add_pPr().append(shading_elm)
                run = p.add_run("\n".join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line.rstrip('\n'))
            continue

        # Handle Tables
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            flush_table(table_lines)
            in_table = False
            table_lines = []

        if not stripped:
            continue

        # Handle Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=0)
            run = h.add_run(stripped[2:])
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            run.font.bold = True
        elif stripped.startswith("## "):
            h = doc.add_heading(level=1)
            run = h.add_run(stripped[3:])
            run.font.color.rgb = RGBColor(0x04, 0x78, 0x57) # Emerald green
            run.font.bold = True
        elif stripped.startswith("### "):
            h = doc.add_heading(level=2)
            run = h.add_run(stripped[4:])
            run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            run.font.bold = True
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            # Check bold in text
            text = stripped[2:]
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="ECFDF5"/>')
            p._p.get_or_add_pPr().append(shading_elm)
            r = p.add_run(stripped[2:])
            r.font.italic = True
            r.font.color.rgb = RGBColor(0x06, 0x5F, 0x46)
        else:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)

    if in_table:
        flush_table(table_lines)

    doc.save(output_path)
    print(f"Successfully generated official Word document: {output_path}")

if __name__ == "__main__":
    create_solution_docx("docs/Solution_Document.md", "AegisFleet_Solution_Document.docx")
